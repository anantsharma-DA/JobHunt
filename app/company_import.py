"""Reads a CSV or Excel (.xlsx) list of company names and careers links."""
import csv
import io
import itertools
import re
import zipfile

from app import normalize

# Limits keep a huge or malicious file from exhausting memory. The API also caps uploads at 5 MB and 2,000 companies.
MAX_ROWS = 2100
MAX_COLUMNS = 20
MAX_CELL_CHARS = 2048
MAX_NAME_CHARS = 200
MAX_UNZIPPED_BYTES = 50 * 1024 * 1024  # an .xlsx is a zip; refuse "zip bombs" that unpack to far more than they look


class FileProblem(ValueError):
    """The uploaded file can't be read; the message is shown to the user."""


def _decode(data):
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise FileProblem("The CSV file's text encoding isn't recognised. Save it as 'CSV UTF-8' and try again.")


def _cell(value):
    return "" if value is None else str(value).strip()[:MAX_CELL_CHARS]


def read_rows(filename, data):
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > MAX_UNZIPPED_BYTES:
                    raise FileProblem("This Excel file is too large once unpacked. Save just the company list as CSV and try again.")
        except zipfile.BadZipFile as exc:
            raise FileProblem("This Excel file couldn't be read. Save it as .xlsx or CSV and try again.") from exc
        try:
            from openpyxl import load_workbook
            from openpyxl.xml import DEFUSEDXML
        except ImportError as exc:
            raise FileProblem("Excel support isn't installed. Close JobHunt and double-click run.bat to install it, or upload a CSV file.") from exc
        if not DEFUSEDXML:
            # Without defusedxml, openpyxl would read XML attacks such as "billion laughs" hidden in an .xlsx file.
            raise FileProblem("Safe Excel reading (defusedxml) isn't installed. Close JobHunt and double-click run.bat to install it, "
                              "or upload a CSV file.")
        try:
            sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
            rows = sheet.iter_rows(values_only=True, max_col=MAX_COLUMNS)
            return [[_cell(value) for value in row] for row in itertools.islice(rows, MAX_ROWS)]
        except FileProblem:
            raise
        except Exception as exc:
            raise FileProblem("This Excel file couldn't be read. Save it as .xlsx or CSV and try again.") from exc
    if name.endswith(".csv"):
        text = _decode(data)
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        try:
            reader = csv.reader(io.StringIO(text), dialect)
            return [[_cell(value) for value in row[:MAX_COLUMNS]] for row in itertools.islice(reader, MAX_ROWS)]
        except csv.Error as exc:
            raise FileProblem("This CSV file couldn't be read. Check it opens in Excel and save it again as CSV.") from exc
    if name.endswith(".xls"):
        raise FileProblem("Old .xls files aren't supported. In Excel use Save As → .xlsx or CSV.")
    raise FileProblem("Upload a .csv or .xlsx file.")


def _header(cell):
    return re.sub(r"[^a-z]+", " ", cell.lower()).strip()


def _is_name_header(h):
    return "compan" in h or h in {"name", "organisation", "organization", "employer", "firm"}


def _is_link_header(h):
    return any(word in h for word in ("link", "url", "career", "website", "site", "page"))


def parse_companies(filename, data):
    """Returns ([{"row", "name", "careers_url"}], [problem messages]).

    A header row is used when one of the first rows names a company column; otherwise column A is the
    company name and column B the careers link.
    """
    rows = read_rows(filename, data)
    name_col, link_col, start = 0, 1, 0
    for i, row in enumerate(rows[:5]):
        heads = [_header(c) for c in row]
        if any(_is_name_header(h) for h in heads):
            name_col = next(j for j, h in enumerate(heads) if _is_name_header(h))
            link_col = next((j for j, h in enumerate(heads) if j != name_col and _is_link_header(h)), None)
            start = i + 1
            break

    companies, problems = [], []
    for row_no, row in enumerate(rows[start:], start=start + 1):
        name = row[name_col].strip()[:MAX_NAME_CHARS] if name_col < len(row) else ""
        link = row[link_col].strip() if link_col is not None and link_col < len(row) else ""
        if not name and not link:
            continue
        if not name:
            problems.append(f"row {row_no}: company name is empty, row skipped")
            continue
        if link and not re.match(r"(?i)^https?://", link):
            if re.match(r"(?i)^[\w-]+(\.[\w-]+)+(/|$)", link):
                link = "https://" + link
            else:
                problems.append(f"row {row_no} ({name}): '{link[:50]}' isn't a web address, added without a link")
                link = ""
        if link and normalize.safe_url(link) is None:
            problems.append(f"row {row_no} ({name}): '{link[:50]}' isn't a usable web address, added without a link")
            link = ""
        companies.append({"row": row_no, "name": name, "careers_url": link or None})
    return companies, problems
